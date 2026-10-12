"""データファイル (CSV / TSV / TXT / Excel) の読み込みとテキスト解析。

pandas (exe で約 10 MB) は使わず、CSV は標準ライブラリ、Excel は openpyxl で直接読む。
"""
from __future__ import annotations

import csv
import io
import os
import re

TABLE_EXTS = (".csv", ".tsv", ".txt", ".dat", ".xlsx", ".xlsm")
EXCEL_EXTS = (".xlsx", ".xlsm")
_ENCODINGS = ("utf-8-sig", "cp932", "latin-1")


class Table:
    """文字列セルの 2 次元表 (行のリスト)。すべての行は同じ列数にそろえる。"""

    def __init__(self, rows: list[list[str]] | None = None):
        rows = rows or []
        width = max((len(r) for r in rows), default=0)
        self.rows = [list(r) + [""] * (width - len(r)) for r in rows]
        self.n_cols = width

    @property
    def empty(self) -> bool:
        return not self.rows or self.n_cols == 0

    @property
    def shape(self) -> tuple[int, int]:
        return len(self.rows), self.n_cols

    def cell(self, r: int, c: int) -> str:
        return self.rows[r][c]

    @property
    def iat(self):  # table.iat[r, c] (pandas と同じ書き方で読めるように)
        table = self

        class _At:
            def __getitem__(self, rc):
                return table.cell(*rc)
        return _At()

    def row(self, r: int) -> list[str]:
        return self.rows[r]

    def column(self, c: int) -> list[str]:
        return [r[c] for r in self.rows]

    def tail(self, start: int) -> "Table":
        return Table(self.rows[start:]) if self.rows[start:] else Table()


def is_number(s) -> bool:
    if s is None:
        return False
    try:
        float(str(s).strip())
        return True
    except ValueError:
        return False


def to_float(s) -> float | None:
    if s is None:
        return None
    text = str(s).strip()
    if not text:
        return None
    try:
        v = float(text)
    except ValueError:
        return None
    return v if v == v and abs(v) != float("inf") else None


def _cell_text(v) -> str:
    if v is None:
        return ""
    return str(v)


def excel_sheets(path: str) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_raw_table(path: str, sheet: str | None = None) -> Table:
    """ヘッダを解釈せず、すべて文字列として読む。空行は除く。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in EXCEL_EXTS:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)  # 数式セルは計算結果の値を読む
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            rows = [[_cell_text(v) for v in r] for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        rows = [r for r in rows if any(c.strip() for c in r)]
        # 右端の空列を落とす
        width = max((max((i + 1 for i, c in enumerate(r) if c.strip()), default=0) for r in rows), default=0)
        return Table([r[:width] for r in rows])
    return parse_text_table(_read_text(path))


def _read_text(path: str) -> str:
    with open(path, "rb") as f:
        raw = f.read()
    for enc in _ENCODINGS:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_text_table(text: str) -> Table:
    """区切り文字 (カンマ / タブ / セミコロン / 空白) を自動判定して表にする。"""
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return Table()
    sample = "\n".join(lines[:50])
    delim = None
    for d in ("\t", ",", ";"):
        if d in sample:
            delim = d
            break
    if delim is None:
        rows = [re.split(r"\s+", ln.strip()) for ln in lines]
    else:
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO("\n".join(lines)), delimiter=delim)]
    return Table(rows)


def detect_header(raw: Table) -> bool:
    """1 行目に数値でないセルがあり、2 行目以降が数値ならヘッダありと判定。"""
    if raw.shape[0] < 2:
        return False
    first = [c for c in raw.row(0) if str(c).strip()]
    second = [c for c in raw.row(1) if str(c).strip()]
    first_numeric = all(is_number(c) for c in first)
    second_numeric = sum(is_number(c) for c in second) >= max(1, len(second) // 2)
    return (not first_numeric) and second_numeric


def split_header(raw: Table, has_header: bool) -> tuple[list[str], Table]:
    if raw.empty:
        return [], raw
    if has_header:
        names = [str(c).strip() or f"列{i + 1}" for i, c in enumerate(raw.row(0))]
        body = raw.tail(1)
    else:
        names = [f"列{i + 1}" for i in range(raw.n_cols)]
        body = raw
    return names, body


def column_values(body: Table, col: int) -> list[float | None]:
    return [to_float(v) for v in body.column(col)]


def parse_rows_text(text: str) -> list[list[float | None]]:
    """貼り付けテキスト (Excel のコピー、カンマ区切りなど) を数値の行列にする。ヘッダ行は除く。"""
    raw = parse_text_table(text)
    if raw.empty:
        return []
    if detect_header(raw):
        raw = raw.tail(1)
    return [[to_float(c) for c in row] for row in raw.rows]


def parse_xy_text(text: str) -> tuple[list[float | None], list[float | None]]:
    """貼り付けテキストの先頭 2 列を x, y として返す。"""
    rows = parse_rows_text(text)
    xs = [r[0] if len(r) > 0 else None for r in rows]
    ys = [r[1] if len(r) > 1 else None for r in rows]
    return xs, ys
